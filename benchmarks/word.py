"""Word documents made from a text, as people's tools write them, for conditions.py: python-docx with each
paragraph's runs split where Word splits them (mid-word, for its spelling and revision marks), a form laid out
as a two-column table, a letterhead and a footer, text boxes (with the copy Word keeps for old readers), tracked
changes, LibreOffice's own .docx, and a scanned page pasted in as a picture."""

import io
import os
import random
import subprocess
import tempfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_LABEL = __import__("re").compile(r"^[-*\s]*\**([A-Z][\w /'()-]{1,40}?)\**\s*:\**\s*(\S.*)$")


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), TYPE


def _runs(p, line, rnd):
    """The line as runs cut at random places, some of them inside a word, as Word cuts them."""
    cuts = sorted(rnd.sample(range(1, len(line)), min(len(line) - 1, rnd.randint(0, 4)))) if len(line) > 1 else []
    for a, b in zip([0] + cuts, cuts + [len(line)]):
        run = p.add_run(line[a:b])
        run.bold = rnd.random() < 0.1


def plain(text, seed):
    rnd = random.Random(seed)
    doc = Document()
    for line in text.split("\n"):
        _runs(doc.add_paragraph(), line, rnd)
    return _save(doc)


def form(text, seed):
    """A form: each "Label: value" line a row of a two-column table, the rest paragraphs."""
    rnd = random.Random(seed)
    doc = Document()
    table = None
    for line in text.split("\n"):
        m = _LABEL.match(line)
        if m:
            if table is None:
                table = doc.add_table(rows=0, cols=2)
            row = table.add_row().cells
            row[0].text = m.group(1)
            row[1].paragraphs[0].text = ""
            _runs(row[1].paragraphs[0], m.group(2), rnd)
        else:
            table = None
            _runs(doc.add_paragraph(), line, rnd)
    return _save(doc)


def header_footer(text, seed):
    """The first line in the page's header, the last two in its footer."""
    rnd = random.Random(seed)
    lines = text.split("\n")
    doc = Document()
    section = doc.sections[0]
    _runs(section.header.paragraphs[0], lines[0], rnd)
    for line in lines[1:-2]:
        _runs(doc.add_paragraph(), line, rnd)
    for line in lines[-2:]:
        _runs(section.footer.add_paragraph(), line, rnd)
    return _save(doc)


_BOX = """<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
  xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
  xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx>
  <w:txbxContent><w:p><w:r><w:t xml:space="preserve">{}</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing>
  </mc:Choice><mc:Fallback><w:pict><v:textbox><w:txbxContent><w:p><w:r><w:t xml:space="preserve">{}</w:t></w:r></w:p>
  </w:txbxContent></v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r>"""


def text_boxes(text, seed):
    """Every third paragraph in a text box, as a callout or a letterhead's address block."""
    from xml.sax.saxutils import escape

    rnd = random.Random(seed)
    doc = Document()
    for n, line in enumerate(text.split("\n")):
        p = doc.add_paragraph()
        if n % 3 == 1 and line.strip():
            p._p.append(etree.fromstring(_BOX.format(escape(line), escape(line))))
        else:
            _runs(p, line, rnd)
    return _save(doc)


def tracked(text, seed):
    """Tracked changes: some words inserted, and a sentence deleted that no longer says anything."""
    rnd = random.Random(seed)
    doc = Document()
    for line in text.split("\n"):
        p = doc.add_paragraph()
        _runs(p, line, rnd)
        if rnd.random() < 0.3 and len(p.runs) > 1:   # the last run as an insertion
            r = p.runs[-1]._r
            ins = etree.SubElement(p._p, qn("w:ins"), {qn("w:id"): str(rnd.randrange(9999)), qn("w:author"): "A"})
            ins.append(r)
        if rnd.random() < 0.2:
            d = etree.SubElement(p._p, qn("w:del"), {qn("w:id"): str(rnd.randrange(9999)), qn("w:author"): "A"})
            r = etree.SubElement(d, qn("w:r"))
            etree.SubElement(r, qn("w:delText")).text = " (an earlier draft's words)"
    return _save(doc)


def libreoffice(text, seed):
    """The text as LibreOffice writes a .docx of it (from HTML: paragraphs and line breaks)."""
    from html import escape

    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, f"doc{seed}.html")
        with open(src, "w") as f:
            f.write("<html><body>" + "".join(f"<p>{escape(x)}</p>" for x in text.split("\n")) + "</body></html>")
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to",
                        "docx:MS Word 2007 XML", "--outdir", tmp, src], capture_output=True, timeout=120, check=True)
        with open(os.path.join(tmp, f"doc{seed}.docx"), "rb") as f:
            return f.read(), TYPE


def pasted_scan(text, seed, page):
    """A scanned page pasted into a Word document as a picture, the document's only content."""
    buf = io.BytesIO()
    page(text).convert("RGB").save(buf, "PNG")
    buf.seek(0)
    doc = Document()
    doc.add_picture(buf, width=Inches(6.5))
    return _save(doc)



def wide_form(text, seed):
    """A form two fields to a row: Label | Value | Label | Value, the labels without their colons."""
    rnd = random.Random(seed)
    doc = Document()
    pending = []

    def flush():
        if not pending:
            return
        table = doc.add_table(rows=0, cols=4)
        for k in range(0, len(pending), 2):
            row = table.add_row().cells
            for c, (label, value) in enumerate(pending[k:k + 2]):
                row[2 * c].text = label
                _runs(row[2 * c + 1].paragraphs[0], value, rnd)
        pending.clear()

    for line in text.split("\n"):
        m = _LABEL.match(line)
        if m:
            pending.append((m.group(1), m.group(2)))
        else:
            flush()
            _runs(doc.add_paragraph(), line, rnd)
    flush()
    return _save(doc)


def tab_form(text, seed):
    """Labels and values lined up with tabs, as people align a form in Word: "Name:<tab><tab>Jane Doe"."""
    rnd = random.Random(seed)
    doc = Document()
    for line in text.split("\n"):
        m = _LABEL.match(line)
        p = doc.add_paragraph()
        if m:
            p.add_run(m.group(1) + ":")
            p.add_run().add_tab()
            if len(m.group(1)) < 12:
                p.add_run().add_tab()
            _runs(p, m.group(2), rnd)
        else:
            _runs(p, line, rnd)
    return _save(doc)
